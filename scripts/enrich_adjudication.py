"""Add question, gold answer, model answer and both reviewers' notes to adjudication.xlsx.

analyze_human_eval.py writes adjudication.xlsx with item ids and the two labels only; this fills in
the text so the adjudicators can decide each disagreement without opening other files.  Existing
final_label / rationale entries are kept.  The key (condition, budget, judge labels) is NOT shown.

Run:  python scripts/enrich_adjudication.py   (then fill final_label + rationale and re-run the analysis)
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
HE = ROOT / "results" / "human_eval_v2"
LABELS = ["CORRECT", "PARTIAL", "INCOMPLETE_NO_ERROR", "INCORRECT", "CORRECT_OTHER_READING", "UNREADABLE"]


def main():
    text, notes = {}, {}
    for f in sorted((HE / "returned").glob("*.xlsx")):
        ws = load_workbook(f, read_only=True)["Items"]
        it = ws.iter_rows(values_only=True)
        head = list(next(it))
        for r in it:
            rec = dict(zip(head, r))
            if rec.get("item_id"):
                text[rec["item_id"]] = (rec["language"], rec["question"], rec["gold_answer"], rec["model_answer"])
                notes[(f.stem, rec["item_id"])] = rec.get("note") or ""
    path = HE / "adjudication.xlsx"
    old = list(load_workbook(path).active.iter_rows(values_only=True))
    head = list(old[0])
    wb = Workbook()
    ws = wb.active
    ws.title = "Adjudication"
    cols = ["item_id", "language", "question", "gold_answer", "model_answer", "reviewer_1", "label_1", "note_1",
            "reviewer_2", "label_2", "note_2", "final_label", "rationale"]
    ws.append(cols)
    for r in old[1:]:
        rec = dict(zip(head, r))
        i = rec["item_id"]
        lang, q, g, m = text[i]
        ws.append([i, lang, q, g, m, rec["reviewer_1"], rec["label_1"], notes.get((rec["reviewer_1"], i), ""),
                   rec["reviewer_2"], rec["label_2"], notes.get((rec["reviewer_2"], i), ""),
                   rec.get("final_label"), rec.get("rationale")])
    for c in ws[1]:
        c.font = Font(bold=True)
    for letter, w in zip("ABCDEFGHIJKLM", [8, 6, 40, 40, 80, 14, 14, 30, 14, 14, 30, 22, 40]):
        ws.column_dimensions[letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"L2:L{ws.max_row}")
    wb.save(path)
    print(f"wrote {path} ({ws.max_row - 1} items)")


if __name__ == "__main__":
    main()
