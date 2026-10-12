"""Copy reviewers' labels (results/human_eval_v2/raw_labels/) into their workbooks in returned/.

The labels were collected by hand and pasted as one table per reviewer; this script puts them into the
same workbooks that were sent out (to_annotate/), so analyze_human_eval.py reads them unchanged.
Checks: every label is a valid choice, every item belongs to that reviewer (assignments_v2.csv),
duplicates agree, and reports what is still missing.  A reviewer listed in SKIP is not written.

Run:  python scripts/fill_returned_workbooks.py
"""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
HE = ROOT / "results" / "human_eval_v2"
RAW, SRC, DST = HE / "raw_labels", HE / "to_annotate", HE / "returned"
LABELS = {"CORRECT", "PARTIAL", "INCOMPLETE_NO_ERROR", "INCORRECT", "CORRECT_OTHER_READING", "UNREADABLE"}

# Reviewer-sheet details.  None = left blank for the reviewer to fill.
REVIEWER = {
    "Chandrahas": {"reviewer_name": "Chandrahas Reddy", "languages_read_well": "English, Hindi, Telugu, Tamil, Kannada",
                   "background (e.g. student, field)": "CSE student", "date_completed": "2026-10-12", "paid? (Y/N)": "N", "ok to be thanked by name in the paper? (Y/N)": "Y"},
    "Akshith": {"reviewer_name": "Akshith Chawla", "languages_read_well": "English, Hindi, Telugu, Kannada",
                "background (e.g. student, field)": "CSE student", "date_completed": "2026-10-12", "paid? (Y/N)": "N", "ok to be thanked by name in the paper? (Y/N)": "Y"},
    "Kapilan": {"reviewer_name": "Kapilan", "languages_read_well": "Tamil, English",
                "background (e.g. student, field)": "CSE student", "date_completed": "2026-10-12", "paid? (Y/N)": "N", "ok to be thanked by name in the paper? (Y/N)": "Y"},
    "Shamik": {"reviewer_name": "Shamik Mukherjee", "languages_read_well": "Bengali, English",
               "background (e.g. student, field)": "CSE student", "date_completed": "2026-10-12", "paid? (Y/N)": "N", "ok to be thanked by name in the paper? (Y/N)": "Y"},
    "TBD_second_Kannada_reader": {"reviewer_name": "Karthik Reddy", "languages_read_well": "Kannada, English",
                                  "background (e.g. student, field)": "CSE student", "date_completed": "2026-10-12", "paid? (Y/N)": "N",
                                  "ok to be thanked by name in the paper? (Y/N)": "Y"},
}
# Second Bengali reader dropped: Bengali items have one reader (Shamik).
SKIP = {"TBD_second_Bengali_reader"}  # dropped by the author 2026-10-12


def read_raw(name: str) -> list[dict]:
    for ext, delim in ((".tsv", "\t"), (".csv", ",")):
        p = RAW / f"{name}{ext}"
        if p.exists():
            with open(p, encoding="utf-8", newline="") as f:
                return list(csv.DictReader(f, delimiter=delim))
    return []


def main():
    assigned: dict[str, set[str]] = {}
    with open(HE / "assignments_v2.csv", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            assigned.setdefault(r["reviewer"], set()).add(r["item_id"])
    DST.mkdir(exist_ok=True)
    for name in sorted(assigned):
        if name in SKIP:
            print(f"{name}: skipped (no named reviewer)")
            continue
        rows, problems, labels = read_raw(name), [], {}
        for r in rows:
            iid, lab = r["item_id"].strip(), r["label"].strip().upper()
            if lab not in LABELS:
                problems.append(f"{iid}: invalid label {lab!r}")
                continue
            if iid not in assigned[name]:
                problems.append(f"{iid}: not assigned to {name}")
                continue
            new = (lab, (r.get("gold_problem") or "").strip().upper(), (r.get("note") or "").strip())
            if iid in labels and labels[iid][0] != lab:
                problems.append(f"{iid}: duplicate with different labels {labels[iid][0]} / {lab}")
            labels[iid] = new
        wb = load_workbook(SRC / f"{name}.xlsx")
        ws = wb["Items"]
        head = [c.value for c in ws[1]]
        ci, cl, cg, cn = (head.index(h) for h in ("item_id", "label", "gold_problem (Y/N)", "note"))
        for row in ws.iter_rows(min_row=2):
            iid = row[ci].value
            if iid in labels:
                lab, gp, note = labels[iid]
                row[cl].value, row[cg].value, row[cn].value = lab, gp or None, note or None
        rv = wb["Reviewer"]
        for row in rv.iter_rows():
            key = row[0].value
            if key in REVIEWER.get(name, {}):
                row[1].value = REVIEWER[name][key]
        wb.save(DST / f"{name}.xlsx")
        missing = sorted(assigned[name] - set(labels))
        print(f"{name}: {len(labels)}/{len(assigned[name])} labelled; missing {len(missing)}"
              + (f": {', '.join(missing)}" if missing else ""))
        for p in problems:
            print("   !", p)


if __name__ == "__main__":
    main()
