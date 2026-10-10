"""Small python-docx helpers for editing the IEEE manuscript in place."""

from __future__ import annotations

import copy
import re

from docx.oxml.ns import qn
from docx.shared import Pt

COND = r"EN-ORIG|EN-MATCH|CS-LAT|ROMAN|NATIVE|MIXED"
TOKEN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\b(?:%s)\b)" % COND)


def _clear(p):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for h in p._element.findall(qn("w:hyperlink")):
        p._element.remove(h)


def add_runs(p, text: str, bold=None, italic=None, size=None):
    """Append text with light markup: **bold**, *italic*, `code`; condition names in small caps."""
    for part in TOKEN.split(text):
        if not part:
            continue
        b, i, sc, code, t = bold, italic, False, False, part
        if part.startswith("`") and part.endswith("`"):
            code, t = True, part[1:-1]
        elif part.startswith("**") and part.endswith("**"):
            b, t = True, part[2:-2]
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            i, t = True, part[1:-1]
        elif re.fullmatch(COND, part):
            sc, t = True, part.lower()
        r = p.add_run(t)
        if b is not None:
            r.bold = b
        if i is not None:
            r.italic = i
        if sc:
            r.font.small_caps = True
        if code:
            r.font.name = "Courier New"
            r.font.size = Pt(7.5)
        if size:
            r.font.size = Pt(size)
    return p


def set_par(p, text: str, bold=None, italic=None, size=None):
    _clear(p)
    return add_runs(p, text, bold=bold, italic=italic, size=size)


def set_caption(p, label: str, text: str):
    """Table captions are 'TABLE N' + line break + caption text in one paragraph."""
    _clear(p)
    p.add_run(label)
    p.add_run().add_break()
    add_runs(p, text)


def insert_after(anchor, template=None):
    """Insert a copy of `template` (default: anchor) after `anchor`; return the new paragraph."""
    from docx.text.paragraph import Paragraph
    src = (template or anchor)._element
    new = copy.deepcopy(src)
    for sp in new.findall(".//" + qn("w:sectPr")):
        sp.getparent().remove(sp)
    anchor._element.addnext(new)
    para = Paragraph(new, anchor._parent)
    _clear(para)
    return para


def delete_par(p):
    p._element.getparent().remove(p._element)


def set_cell(cell, text: str, bold=None):
    p = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        delete_par(extra)
    # keep formatting of the first run if present
    fmt = None
    if p.runs:
        r0 = p.runs[0]
        fmt = (r0.bold, r0.font.size, r0.font.small_caps)
    _clear(p)
    add_runs(p, text, bold=bold if bold is not None else (fmt[0] if fmt else None),
             size=(fmt[1].pt if fmt and fmt[1] else None))


def fill_table(table, rows: list[list[str]], header: list[str] | None = None):
    """Resize table body to len(rows) (copying the last row) and write all cells."""
    start = 1 if header is not None else 0
    if header is not None:
        for c, h in zip(table.rows[0].cells, header):
            set_cell(c, h)
    body = table.rows[start:]
    while len(table.rows) - start < len(rows):
        table._tbl.append(copy.deepcopy(table.rows[-1]._tr))
    while len(table.rows) - start > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for r, vals in zip(table.rows[start:], rows):
        for c, v in zip(r.cells, vals):
            set_cell(c, v)


def add_column(table, header: str, values: list[str], width_from: int = -1):
    """Append a column by copying the cells of an existing column."""
    for i, row in enumerate(table.rows):
        tc = row.cells[width_from]._tc
        new = copy.deepcopy(tc)
        tc.addnext(new) if width_from != -1 else row._tr.append(new)
    grid = table._tbl.tblGrid
    grid.append(copy.deepcopy(grid.gridCol_lst[width_from]))
    from docx.table import _Cell
    for i, row in enumerate(table.rows):
        cell = row.cells[-1]
        set_cell(cell, header if i == 0 else values[i - 1])


def copy_table_after(anchor_par, table):
    new = copy.deepcopy(table._tbl)
    anchor_par._element.addnext(new)
    from docx.table import Table
    return Table(new, anchor_par._parent)
