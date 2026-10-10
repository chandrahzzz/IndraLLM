"""Build the 10-page revised IEEE manuscript (128 / 512 / 1,024-token results, two judges).

Starts from paper_v2/source/IndraLLM_IEEE_FINAL.docx (preserving its styles and section
layout), rewrites the text, rebuilds the tables from the JSON result files, replaces the
figures, and writes paper_v2/revised/IndraLLM_IEEE_REVISED.docx.

Run after: make_all.py (both label sets), analyze_exp004.py, analyze_budgets.py,
make_figures.py, make_fig1.py.
"""

from __future__ import annotations

import copy
import json
import sys
from pathlib import Path

import docx

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_revised_docx as base  # noqa: E402  (text, provisions, number formatting)
from docx_helpers import (  # noqa: E402
    add_column, add_runs, copy_table_after, delete_par, fill_table, insert_after, set_caption, set_par,
)

ROOT = base.ROOT
NO, NJ, NE, L6 = base.NO, base.NJ, base.NE, base.L6
NB = json.loads((ROOT / "paper_v2/output/numbers_budgets.json").read_text(encoding="utf-8"))
fp, sgn, neg, by_cmp, set_col_widths = base.fp, base.sgn, base.neg, base.by_cmp, base.set_col_widths
OUTDOC = ROOT / "paper_v2/revised/IndraLLM_IEEE_REVISED.docx"
TEXT = dict(base.TEXT)
G3 = ["EN-MATCH", "CS-LAT", "MIXED"]
BUD = ["128", "512", "1024"]

# --------------------------------------------------------------------------------------
# Text (overrides of the 13-page draft)
# --------------------------------------------------------------------------------------
TEXT["abstract"] = (
    "Comparisons of LLM factual accuracy across languages and scripts usually change the question along with "
    "its representation. We hold 100 propositions on 20 Indian public-policy topics fixed and render each in "
    "English, native script, Romanized text, Latin-script code-switching (CS-LAT) and mixed-script "
    "code-switching (MIXED) in five Indic languages (Hindi, Bengali, Tamil, Telugu, Kannada), and ask how "
    "much the measured differences depend on two evaluation choices: the output-token budget and the judge. "
    "With Qwen3.8-27B and a 128-token cap, judge-assessed accuracy is 60% for a template-matched English "
    "prompt and 49%, 37%, 32% and 26% for CS-LAT, Romanized, native-script and MIXED prompts, and MIXED is "
    "23 points below CS-LAT (p < 0.001; topic-level p = 0.004). Almost every response hit that cap, and "
    "Brahmic-script answers fit half as many characters into it. Re-answering the key conditions with 512 "
    "and 1,024 tokens removes the MIXED–CS-LAT gap under both a self-judge (+4 and −3 points; p ≥ 0.57) and "
    "a judge from another model family (−6 and −5; p ≥ 0.34). The English–CS-LAT gap behaves differently: under the self-judge it "
    "grows with the budget, to 24 points at 1,024 tokens (topic-level p = 0.018), while under the second "
    "judge it is 4–12 points and never significant. The two judges agree on 75–81% of outputs "
    "(κ = 0.47–0.61), and no human validation of either was performed. Measured representation effects are "
    "therefore sensitive to the output budget and to the evaluator, and should not be read as effects of "
    "representation without checking both. Prompts, outputs, judgements, a gold-answer ledger and code are "
    "released.")

TEXT["index"] = ("multilingual evaluation, code-switching, script, Indic languages, output budget, "
                 "LLM-as-a-judge, judge agreement, paired evaluation")

TEXT[9] = (
    "We build the comparison so that they can be. Each of 100 propositions is rendered in five paired "
    "representations (Table I). Two of them, CS-LAT and MIXED, are meant to differ only in whether the Indic "
    "words are written in Latin or Brahmic script; this holds word for word in Hindi, Bengali and Tamil, "
    "while in Telugu and Kannada the question word also differs (Section III-A). We add a second English "
    "condition after finding that the original English template was more specific than the template shared "
    "by the other four conditions (Section III-B). We then ask whether the differences we measure survive "
    "two changes to the evaluation that representation studies rarely make: a larger output budget (512 and "
    "1,024 tokens instead of 128) and a judge from a different model family.")

TEXT[13] = (
    "**Research questions.** RQ1: at a fixed output budget and judge, does judge-assessed accuracy depend on "
    "representation when the proposition is held fixed? RQ2: does mixed-script code-switching lower accuracy "
    "relative to the same code-switched words in Latin script? RQ3: do these differences change when the "
    "output budget grows? RQ4: do they change with the judge, and how far do two judges agree? RQ5: is the "
    "script contrast associated with input tokenization length? Hallucination detectors and mitigation are "
    "outside the scope of this paper.")

TEXT[14] = (
    "**Contributions.** (1) A paired six-condition design for Indic factual QA, with a matched-English "
    "control that removes a template confound we found in our first run. (2) A controlled three-budget "
    "rerun showing that the 23-point mixed-script penalty measured at 128 tokens disappears at 512 and 1,024 "
    "tokens under two judges, while the English–code-switching gap moves in the opposite direction under one "
    "judge only. (3) Judge disagreement as a result: harmonising the judge configuration shrinks the "
    "English–code-switching gap at 128 tokens from 17 to 11 points, the two judges agree on 75–81% of outputs, "
    "and whether the English advantage is significant depends on which judge is used. (4) Inference that "
    "treats the 20 topics as the unit of independent evidence, and a release of prompts, outputs, both "
    "judges' labels, a gold-answer ledger and evaluation scripts.")

TEXT[15] = (
    "**Findings in brief.** Evaluating multilingual LLMs across representation conditions is sensitive to the "
    "output-token budget and to the choice of evaluator, and observed differences should not automatically be "
    "interpreted as representation effects. At 128 tokens, accuracy falls from English to the Indic "
    "representations and mixed script is lowest. Almost every response hits that cap and Brahmic-script "
    "answers are about half as long in characters; with 512 or 1,024 tokens, MIXED and CS-LAT do not differ "
    "under either judge. The English advantage over Latin-script code-switching is significant at the topic "
    "level only under the self-judge at 1,024 tokens; under the cross-family judge it is small at every "
    "budget. Neither judge has been validated against human labels, so we report this as judge sensitivity, "
    "not as evidence that either judge is right.")

TEXT[20] = (
    "**Factuality evaluation and LLM judges.** We use closed-book questions with a gold answer, closer in "
    "spirit to TruthfulQA-style checks than to long-form factual precision scoring [11], [12]. LLM graders "
    "are convenient but biased, for example towards their own outputs, answer position and length [13], "
    "[24], [25], and their agreement with humans and with each other is lower in many non-English languages, "
    "including Indic languages [26], [27]. We therefore grade with two judges from different model families "
    "and report their disagreement as a result.")

TEXT["fig1"] = (
    "Fig. 1. Study design. Each proposition is rendered in six conditions and answered once by Qwen3.8-27B "
    "under greedy decoding with a 128-token limit; responses are graded against the gold answer by the same "
    "model (J-QWEN). The follow-up (Section VI-B) re-answers three conditions with 512 and 1,024-token limits "
    "and grades every budget with J-QWEN and with gpt-oss-120b (J-OSS). Key contrasts: EN-MATCH vs. CS-LAT "
    "(code-switching) and CS-LAT vs. MIXED (script).")

TEXT[28] = base.TEXT[28].replace(
    "Fluent speakers of all five languages", "Fluent speakers of all five languages").replace(
    "No fluent speaker has yet rated the naturalness or equivalence of these templates; the templates and a "
    "review protocol are released.",
    "Fluent speakers of all five languages, including the author (Telugu) and a teammate (Hindi, Telugu, "
    "Kannada), reviewed the prompt templates and judged them natural and correct. This was a yes/no check "
    "without item-level ratings; the reviewers saw the templates, not the model's outputs, and the check "
    "establishes neither full semantic equivalence nor the validity of the judges.")

TEXT[34] = base.TEXT[34].replace(
    "The author and, separately, volunteer reviewers known to the author checked the gold answers for factual "
    "correctness against public sources. That check was not logged item by item, the reviewers were not "
    "trained annotators, and no agreement statistic was computed, so we report it as a check rather than a "
    "measurement. For this version we also matched each gold answer to its governing provision and reviewed "
    "it for currency and for completeness relative to the question actually asked (Appendix A, Table VIII; "
    "a per-topic ledger is released).",
    "The author and, separately, volunteer reviewers checked the gold answers for factual correctness against "
    "public sources, and the author verified each answer against its governing provision, as in force in "
    "October 2026, for currency and for completeness relative to the question actually asked (Appendix A, "
    "Table VII; a per-topic ledger is released). The volunteer check was not logged item by item and no "
    "agreement statistic was computed, so we report it as a check rather than a measurement.").replace(
    ", and are not used here.", "; they are not used.")

TEXT[37] = base.TEXT[37].replace(
    "in the follow-up, 84–85% of 512-token responses begin with exactly the 128-token response to the same "
    "prompt, and the rest diverge.",
    "84–85% of 512-token responses begin with exactly the 128-token response to the same prompt, and 78–91% "
    "of 1,024-token responses with the 512-token response; the rest diverge.")

TEXT[38] = base.TEXT[38].replace(
    "The provider's finish reason was not logged in the main runs; it was logged in the follow-up.",
    "The provider's finish reason was not logged in the main runs; it was logged in both follow-up runs.")

TEXT[39] = base.TEXT[39].replace(
    "No judge has been validated against human labels on these outputs.",
    "No judge has been validated against human labels on these outputs, and we do not treat either as the "
    "more accurate one.")

TEXT[40] = base.TEXT[40].replace(
    "Logistic GEE with independence working correlation and robust covariance clustered at prompt, "
    "proposition and topic level [18] is reported for comparison;",
    "Logistic GEE with independence working correlation and robust covariance clustered at prompt, "
    "proposition and topic level [18] is reported in Appendix B;").replace(
    "The follow-up (Section VI-B) has three contrasts, fixed in the analysis script before its outputs were "
    "graded: CS-LAT vs. EN-MATCH and MIXED vs. CS-LAT at 512 tokens, and the change in each gap from 128 to "
    "512 tokens.",
    "The follow-up (Section VI-B) has the same three contrasts at every budget, fixed in the analysis scripts "
    "before its outputs were graded: CS-LAT vs. EN-MATCH, MIXED vs. CS-LAT and MIXED vs. EN-MATCH, plus the "
    "change in each gap between budgets, each under both judges.")

TEXT[43] = base.TEXT[43].replace("(Table IX).", "(Table VIII).")
TEXT["tab2"] = base.TEXT["tab2"].replace("text and Table IX", "text and Table VIII")

TEXT[59] = base.TEXT[59].replace("Table III and Fig. 4a show what this does.",
                                 "Fig. 4a and Table IX show what this does.").replace(
    "The English-vs-Indic differences are consistent in direction, but at the level of topics only the larger "
    "ones approach significance, and the English–CS-LAT gap does not.",
    "At this budget the English-vs-Indic differences are consistent in direction, but at the level of topics "
    "only the larger ones approach significance, and the English–CS-LAT gap does not.")

TEXT[65] = base.TEXT[65].replace("Accuracy by language is in Fig. 3 and Appendix Table X.",
                                 "Accuracy by language is in Fig. 3.")

TEXT[76] = (
    "To separate script from budget we re-ran EN-MATCH, CS-LAT and MIXED (300 prompts) twice, with 512 and "
    "with 1,024-token limits, changing nothing else: same model identifier, system prompt, decoding settings "
    "and prompts. Every budget was graded by J-QWEN and by J-OSS, so each comparison across budgets uses one "
    "judge (Table III, Fig. 5b). At 1,024 tokens the cap no longer binds for English (95 of 100 responses "
    "stop on their own) and binds less for code-switched answers (61 CS-LAT and 36 MIXED responses stop). "
    "The MIXED–CS-LAT gap, −23 points at 128 tokens under J-QWEN, is +4 at 512 and −3 at 1,024 "
    "(p = 0.57 and 0.70); under J-OSS it is −8, −6 and −5, never significant. Between 128 and 1,024 tokens it "
    "narrows by 20 points under J-QWEN (topic-bootstrap CI [5, 34], permutation p = 0.023). On the 60 "
    "word-identical propositions the picture is the same (J-QWEN at 1,024 tokens: 65% vs. 60%, p = 0.58). "
    "The English–CS-LAT gap moves the other way under J-QWEN, from −11 at 128 to −21 at 512 and −24 at 1,024 "
    "tokens (exact p < 0.001; topic-level p = 0.018, Holm-adjusted 0.036 within the three follow-up "
    "contrasts), whereas under J-OSS it is −12, −4 and −7 and not significant at any budget (p ≥ 0.050).")

TEXT["tab3"] = (
    "Output-budget follow-up: the same 300 prompts answered with 128, 512 and 1,024 tokens, each budget "
    "graded by both judges. Accuracy (%) per condition; gaps in points with exact McNemar p / topic-level "
    "permutation p. Cap: responses (of 100) that hit the output limit, EN-MATCH / CS-LAT / MIXED.")

TEXT[77] = (
    "Three conclusions follow. First, the 128-token MIXED penalty is not evidence that script alternation "
    "harms recall: once the budget allows longer answers, mixed-script and Latin-script code-switched prompts "
    "yield similar accuracy under both judges. Its most likely source is that Brahmic-script answers reach "
    "the cap before reaching the fact, and answers a human would call incomplete are then graded as wrong; "
    "the EN-MATCH answer on the MTP Act, for instance, correctly states the one-practitioner rule for the "
    "first 20 weeks, is cut off before the 20–24-week rule, and is graded HALLUCINATED. Second, NATIVE and "
    "MIXED behave alike at 128 tokens (32% and 26%, p = 0.29) although MIXED has five script boundaries per "
    "prompt and NATIVE one, which also argues against a boundary-density account. Third, the English "
    "advantage is unstable: it grows with the budget under one judge and stays small under the other. We "
    "cannot tell from these data which judge is closer to a careful human reader, so the instability itself "
    "is the finding.")

TEXT[79] = (
    "The judge matters in three ways. *Configuration.* Our first run graded the Indic conditions at "
    "temperature 0.3 with the answering system prompt and EN-MATCH at temperature 0 with an evaluator prompt. "
    "Harmonising the configuration changed 6–10% of the labels in each re-graded condition and moved CS-LAT "
    "accuracy from 43% to 49%, so the stored English–CS-LAT gap (17 points) exceeded the harmonised one (11). "
    "*Family.* Table IV gives agreement between J-QWEN and J-OSS by budget and condition. Agreement is 74.7% "
    "at 128 tokens (κ = 0.47), 80.0% at 512 (κ = 0.59) and 81.3% at 1,024 (κ = 0.61), with wide "
    "topic-bootstrap intervals for κ. Disagreements run mostly one way: J-QWEN labels CORRECT 30–36% of the "
    "answers J-OSS labels wrong, whereas J-OSS labels CORRECT 6–11% of the answers J-QWEN labels wrong. In "
    "the disagreements we inspected, J-QWEN accepted answers that were on topic and accurate but did not state "
    "the gold fact (for example, answers that only name the governing Act), and J-OSS rejected them; the "
    "rubric (“judge only factual correctness and whether it actually addresses the question”) "
    "does not say which reading is intended, and truncated answers make such cases common. *Conclusions.* "
    "The MIXED–CS-LAT conclusion depends on the judge at 128 tokens (significant under J-QWEN only) but not "
    "at 512 or 1,024 tokens (not significant under either). The English–CS-LAT conclusion depends on the "
    "judge at both larger budgets. No human gold-standard labels were collected for any output, so these "
    "results establish that conclusions are sensitive to the judge, not which judge is right.")

TEXT["tab4"] = (
    "Agreement between the two judges on the same 300 outputs at each budget. Agree: share of identical "
    "labels; κ: Cohen's kappa with topic-bootstrap 95% CI; Q✓O✗: J-QWEN CORRECT and J-OSS wrong; Q✗O✓: the "
    "reverse. No human labels were collected.")

TEXT[80] = base.TEXT[80].replace("Table VI turns", "Table V turns")
TEXT["tab5"] = base.TEXT["tab6"]
TEXT["tab6"] = "Claims examined and not supported, or supported only under some evaluation settings."

TEXT[83] = (
    "What the data support is narrower than the design promised, and different from what our first run "
    "suggested. Evaluating multilingual LLMs across representation conditions is sensitive to the "
    "output-token budget and to the choice of evaluator, and observed differences should not automatically "
    "be interpreted as representation effects. For this model and these 20 topics, the clearest "
    "representation effect at a 128-token budget, the mixed-script penalty, disappears when answers can be "
    "longer, under both judges; the English advantage over code-switching appears or disappears with the "
    "judge. We do not claim that native script or code-switching damages factual recall, and we make no "
    "claim about other models.")

TEXT[84] = (
    "The practical lesson is methodological. Output caps are often set low to save cost, and tokenizers make "
    "some scripts more expensive per character, so a shared cap gives answers in those scripts less room. "
    "Studies that compare languages or scripts through generated answers should report how often the cap "
    "binds in each condition and show that conclusions hold at a non-binding budget; judges should be "
    "configured identically across conditions, at least two judges from different families should be used, "
    "and their agreement should be reported per condition. Whether this cautionary result is novel enough "
    "for a given venue depends on prior work on tokenization premiums and multilingual judge reliability "
    "[5], [6], [26], [27], which documents the ingredients but, to our knowledge, not their combined effect "
    "on paired representation contrasts. Table VI lists claims we examined.")

TEXT[87] = (
    "**Topics.** Twenty topics with heavy within-topic concentration bound what any statistic can show "
    "(effective n ≈ 67 of 500). A 25-proposition expansion has been written but not evaluated. **Output "
    "budget.** At 1,024 tokens, 39 CS-LAT and 64 MIXED responses still hit the cap; we re-ran three of the six "
    "conditions and did not test an instruction fixing the answer language. **Judges.** Two LLM judges that "
    "disagree on a fifth to a quarter of items, and no human labels on any output; a stratified, blinded "
    "sample of 210 outputs and an annotation protocol are released but not annotated. **Questions and gold "
    "answers.** One generic template that under-specifies the target for several topics; volunteer gold "
    "checks without an item-level record; one superseded scheme. **Prompts.** The fluent-speaker review was "
    "an informal yes/no check of the templates; CS-LAT and MIXED differ by one word in Telugu and Kannada, "
    "and Hindi ROMAN is not a transliteration of Hindi NATIVE. **Models and languages.** One informative "
    "model, served under an API identifier whose underlying weights we cannot independently verify, and five "
    "languages. **Single run.** No repeated sampling; hosted greedy decoding was not fully deterministic, and "
    "API behaviour can change over time.")

TEXT[91] = (
    "All numbers, tables and figures are produced from the raw files by the scripts in `paper_v2/scripts/` "
    "(Appendix E). The two follow-up runs cost $1.10 (512 tokens plus re-grading of the 128-token outputs) "
    "and $1.23 (1,024 tokens) at the provider's list prices in October 2026; all 2,700 follow-up judgements "
    "parsed, and no call failed after retries. The project's budget ledger records $0.21 for the earlier "
    "runs but priced every call at a flat $0.08 per million tokens; at list prices the Qwen3.8-27B "
    "generation and grading in EXP-002 and EXP-003 would cost about $2.1. The provider's billing record is "
    "the authoritative figure.")

TEXT[94] = base.TEXT[94].replace("Table VIII lists", "Table VII lists").replace(
    "the provision against which each gold answer was reviewed, the outcome of that review,",
    "the provision against which each gold answer was verified, the outcome,")
TEXT["tab7"] = base.TEXT["tab9"].replace("used to review the gold answer, review outcome",
                                         "used to verify the gold answer, verification outcome")
TEXT["appB"] = base.TEXT["appB"] + " Table VIII gives all pairwise contrasts and Table IX the GEE estimates."
TEXT["tab8"] = base.TEXT["tab10"]
TEXT["tab9"] = base.TEXT["tab3"]

TEXT["fig5"] = (
    "Fig. 5. (a) Mean response length for EN-MATCH, CS-LAT and MIXED at 128, 512 and 1,024 tokens; labels give "
    "the % of responses that hit the cap. (b) Accuracy of the same 300 prompts at each budget, graded by "
    "J-QWEN (solid) and J-OSS (dashed). (c) Condition-level prompt characters per token (Qwen3.8 tokenizer) "
    "against accuracy at 128 tokens (six points; descriptive only).")

TEXT[110] = base.TEXT[110].replace(
    "All 1,800 follow-up judgements were parsed.",
    "All 2,700 follow-up judgements (1,800 at 128/512 tokens, 900 at 1,024 tokens) were parsed.")

TEXT[114] = (
    "Python 3.11 with numpy, pandas, scipy, statsmodels, tokenizers and matplotlib. From the repository root, "
    "`python paper_v2/scripts/make_all.py` (stored labels) and `make_all.py --labels jqwen` read the two "
    "prediction files; `scripts/run_exp004_budget_rejudge.py` and `scripts/run_exp005_budget1024.py` "
    "(`--stage gen|judge`; Groq key required) produced the follow-ups, with finish reasons, failures and a "
    "per-call spend log in `results/EXP-004-budget-rejudge/` and `results/EXP-005-budget1024/`; "
    "`analyze_exp004.py`, `analyze_budgets.py` and `make_figures.py` write the remaining numbers and figures. "
    "The tokenizer check downloads the Qwen/Qwen2.5-7B and Qwen/Qwen3.8-27B tokenizers. Sign-flip tests "
    "enumerate all 2²⁰ patterns; bootstraps and the language permutation use fixed seeds. Topic-level p "
    "values in an earlier draft were Monte Carlo approximations and differ from the exact values here in the "
    "third decimal. That draft's language-dependence p value (0.15) and numeral check (60% agreement) could "
    "not be reproduced because the script that produced them was not in the repository; they are replaced "
    "by the procedures described above.")

NEW_REFS = base.NEW_REFS[:]
NEW_REFS[0] = ("J. A. Husain *et al.*, “RomanSetu: Efficiently unlocking multilingual capabilities of large "
               "language models via Romanization,” in *Proceedings of the 62nd annual meeting of the "
               "association for computational linguistics*, 2024.")
NEW_REFS += [
    "I. Watts, V. Gumma, A. Yadavalli, V. Seshadri, M. Swaminathan, and S. Sitaram, “PARIKSHA: A "
    "large-scale investigation of human-LLM evaluator agreement on multilingual and multi-cultural data,” "
    "in *Proceedings of the 2024 conference on empirical methods in natural language processing*, 2024, "
    "pp. 7900–7932.",
    "X. Fu and W. Liu, “How reliable is multilingual LLM-as-a-judge?” in *Findings of the "
    "association for computational linguistics: EMNLP 2025*, 2025, pp. 11040–11053.",
]


for key, must in [(28, "templates, not the model"), (34, "October 2026"), (37, "78–91%"),
                  (38, "both follow-up runs"), (39, "more accurate one"), (40, "same three contrasts"),
                  (40, "Appendix B;"), (43, "(Table VIII)"), ("tab2", "Table VIII"), (59, "Table IX show"),
                  (65, "in Fig. 3."), (80, "Table V turns"), (94, "Table VII lists"), (110, "2,700")]:
    assert must in TEXT[key], f"text replacement failed for {key}: {must}"


def pfmt(p):
    return "<.001" if p < 0.001 else fp(p).lstrip("0")


def main():
    d = docx.Document(base.SRC)
    P = list(d.paragraphs)
    T = list(d.tables)

    # Templates copied before anything is modified.
    sect2_end = copy.deepcopy(P[44]._element)   # ends a two-column section
    sect1_end = copy.deepcopy(P[46]._element)   # ends a one-column (full-width) section
    full_caption = P[102]                         # caption style used for full-width tables
    tbl_full = copy.deepcopy(T[6]._tbl)           # 7-column full-width table
    tbl_small = T[3]                               # 4-column one-column table

    # ---------------- text ----------------
    set_par(P[0], TEXT["title"])
    set_par(P[4], "")
    add_runs(P[4], "Abstract—", bold=True, italic=True)
    add_runs(P[4], TEXT["abstract"], bold=True)
    set_par(P[5], "")
    add_runs(P[5], "Index Terms—", bold=True, italic=True)
    add_runs(P[5], TEXT["index"], bold=True)
    for i in [9, 13, 14, 15, 18, 19, 20, 28, 30, 34, 36, 37, 38, 39, 40, 41, 43, 53, 54, 55, 57, 59, 65, 66,
              67, 68, 69, 74, 75, 76, 77, 78, 79, 80, 83, 84, 87, 89, 91, 94, 110, 112, 114]:
        set_par(P[i], TEXT[i])
    for i, key in [(23, "fig1"), (49, "fig2"), (52, "fig3"), (62, "fig4"), (72, "fig5")]:
        set_par(P[i], TEXT[key])
    set_caption(P[11], "TABLE I", TEXT["tab1"])

    # ---------------- Table II: accuracy (J-QWEN, 128 tokens) ----------------
    set_caption(P[45], "TABLE II", TEXT["tab2"])
    rows = []
    for c in L6:
        a = NJ["accuracy"][c]
        k = a["k"]
        wl, wh = a["wilson"]
        bl, bh = a["topic_boot_ci"]
        st, oss = NO["accuracy"][c]["k"], NE["orig128"]["J-OSS"]["accuracy"][c]
        if c == "EN-MATCH":
            row = [c, f"{k:.1f} ({k}/100)", f"[{wl:.1f}, {wh:.1f}]", f"[{bl:.0f}, {bh:.0f}]", "ref.", "–", "–", "–"]
        elif c == "EN-ORIG":
            x = by_cmp(NJ, "EN-MATCH", "EN-ORIG")
            row = [c, f"{k:.1f} ({k}/100)", f"[{wl:.1f}, {wh:.1f}]", f"[{bl:.0f}, {bh:.0f}]", f"{sgn(k - 60)}†",
                   fp(x["p_exact"]), "–", fp(x["topic_perm_p"])]
        else:
            x = by_cmp(NJ, c, "EN-MATCH")
            row = [c, f"{k:.1f} ({k}/100)", f"[{wl:.1f}, {wh:.1f}]", f"[{bl:.0f}, {bh:.0f}]", sgn(k - 60),
                   fp(x["p_exact"]), fp(x["holm_exact"]), f"{fp(x['topic_perm_p'])} ({fp(x['holm_topic'])})"]
        rows.append(row + [str(st), f"{oss:.0f}"])
    add_column(T[1], "Stored", [""] * 6)
    add_column(T[1], "J-OSS", [""] * 6)
    fill_table(T[1], rows, header=["Condition", "Acc. (%)", "Wilson CI", "Topic-boot. CI", "Δ", "Exact p",
                                   "Holm p", "Topic p (Holm)", "Stored", "J-OSS"])
    set_col_widths(T[1], [0.75, 0.85, 0.85, 0.75, 0.45, 0.8, 0.8, 0.95, 0.5, 0.5])

    # ---------------- Table III: budgets (full width, after P[76]) ----------------
    cap3 = insert_after(P[76], full_caption)
    set_caption(cap3, "TABLE III", TEXT["tab3"])
    cap3._element.addprevious(copy.deepcopy(sect2_end))
    cap3._element.addnext(tbl_full)
    from docx.table import Table
    t3 = Table(tbl_full, cap3._parent)
    tbl_full.addnext(copy.deepcopy(sect1_end))
    add_column(t3, "", [""] * (len(t3.rows) - 1))
    rows = []
    for j in ["J-QWEN", "J-OSS"]:
        for b in BUD:
            acc = NB["accuracy"][j][b]
            con = {f"{c['cmp']}-{c['ref']}": c for c in NB["contrasts"][j][b]}
            ce, mc = con["CS-LAT-EN-MATCH"], con["MIXED-CS-LAT"]
            capped = " / ".join(str(100 - NB["generation"][b][c]["stopped_before_cap"]) for c in G3)
            rows.append([j, f"{int(b):,}"] + [f"{acc[c]['pct']:.0f}" for c in G3] +
                        [f"{sgn(ce['delta_pp'])} ({pfmt(ce['p_exact'])} / {pfmt(ce['topic_perm_p'])})",
                         f"{sgn(mc['delta_pp'])} ({pfmt(mc['p_exact'])} / {pfmt(mc['topic_perm_p'])})", capped])
    fill_table(t3, rows, header=["Judge", "Budget", "EN-MATCH", "CS-LAT", "MIXED", "CS-LAT − EN-MATCH",
                                 "MIXED − CS-LAT", "Cap (E/C/M)"])
    set_col_widths(t3, [0.7, 0.6, 0.75, 0.65, 0.65, 1.3, 1.3, 0.95])

    # ---------------- Table IV: judge agreement (one column, after P[79]) ----------------
    cap4 = insert_after(P[79], P[81])
    set_caption(cap4, "TABLE IV", TEXT["tab4"])
    t4 = copy_table_after(cap4, tbl_small)
    add_column(t4, "", [""] * 5)
    add_column(t4, "", [""] * 5)
    rows = []
    for b in BUD:
        for c in G3 + ["ALL"]:
            x = NB["judge_agreement"][b][c]
            lo, hi = x["kappa_topic_ci"]
            rows.append([f"{int(b):,}" if c == "EN-MATCH" else "", "All" if c == "ALL" else c,
                         f"{x['agreement']:.2f}", f"{x['kappa']:.2f} [{lo:.2f}, {hi:.2f}]",
                         str(x["qwen_correct_oss_wrong"]), str(x["qwen_wrong_oss_correct"])])
    fill_table(t4, rows, header=["Budget", "Condition", "Agree", "κ [95% CI]", "Q✓O✗", "Q✗O✓"])
    set_col_widths(t4, [0.5, 0.75, 0.45, 1.0, 0.45, 0.45])

    # ---------------- Table V: tipping points ----------------
    set_caption(P[81], "TABLE V", TEXT["tab5"])
    tp = NJ["tipping"]
    fill_table(T[3], [[f, f"{tp['CS-LAT_vs_EN-MATCH'][f]:.2f}", f"{tp['MIXED_vs_EN-MATCH'][f]:.2f}",
                       f"{tp['MIXED_vs_CS-LAT'][f]:.2f}"] for f in ["0.00", "0.04", "0.08", "0.12", "0.16"]],
               header=["FPR", "CS-LAT vs. EN-MATCH", "MIXED vs. EN-MATCH", "MIXED vs. CS-LAT"])

    # ---------------- Table VI: claims ----------------
    set_caption(P[85], "TABLE VI", TEXT["tab6"])
    fill_table(T[4], [
        ["Script alternation (CS-LAT → MIXED) lowers recall",
         "Not supported: only at 128 tokens; +4 (512), −3 (1,024), n.s. under both judges"],
        ["English advantage over CS-LAT",
         "Judge-dependent: topic p = 0.018 at 1,024 tokens (J-QWEN); n.s. at every budget (J-OSS)"],
        ["Prompt token fertility mediates CS-LAT → MIXED", "Not detected (p = 0.08, wrong sign; exploratory)"],
        ["Language dependence of the profile", "Not detected (p = 0.17), low power"],
        ["Rank agreement with ALLaM-2-7B", "Uninformative (floor; ρ = 0.67, p = 0.22)"],
        ["Errors rise with script transitions", "Not supported (NATIVE ≈ MIXED)"],
        ["Original 20-fold truncation ratio", "Not reproducible from files"],
    ], header=["Examined claim", "Outcome"])

    # ---------------- Appendix A: Table VII topics ----------------
    set_caption(P[96], "TABLE VII", TEXT["tab7"])
    add_column(T[5], "", [""] * 20)
    add_column(T[5], "", [""] * 20)
    topics = sorted(NJ["by_topic"], key=lambda t: (-NJ["by_topic"][t]["EN-MATCH"], -NJ["by_topic"][t]["EN-ORIG"], t))
    fill_table(T[5], [[t, base.PROVISION[t][0], base.PROVISION[t][1]] + [str(NJ["by_topic"][t][c]) for c in L6]
                      for t in topics], header=["Topic", "Provision verified", "Gold"] + L6)
    set_col_widths(T[5], [1.75, 1.95, 0.45, 0.5, 0.55, 0.45, 0.45, 0.45, 0.45])

    # ---------------- Appendix B: text, Table VIII pairwise, Table IX GEE ----------------
    pb = insert_after(P[100], P[99])
    set_par(pb, TEXT["appB"])
    set_caption(P[102], "TABLE VIII", TEXT["tab8"])
    order = [("CS-LAT", "EN-MATCH"), ("ROMAN", "EN-MATCH"), ("NATIVE", "EN-MATCH"), ("MIXED", "EN-MATCH"),
             ("MIXED", "CS-LAT"), ("EN-MATCH", "EN-ORIG"), ("MIXED", "ROMAN"), ("MIXED", "NATIVE"),
             ("ROMAN", "CS-LAT"), ("NATIVE", "CS-LAT"), ("NATIVE", "ROMAN")]
    rows = []
    for cmp_, ref in order:
        x = by_cmp(NJ, cmp_, ref)
        lo, hi = x["topic_boot_ci_pp"]
        rows.append([f"{cmp_} vs. {ref}", sgn(x["delta_pp"]), str(x["b"]), str(x["c"]), fp(x["p_exact"]),
                     fp(x["topic_perm_p"]), neg(f"[{lo:.0f}, {hi:.0f}]")])
    fill_table(T[6], rows, header=["Contrast", "Δ (pp)", "b", "c", "Exact p", "Topic p", "Topic 95% CI (pp)"])

    set_caption(P[56], "TABLE IX", TEXT["tab9"])
    g = NJ["gee"]
    rows = []
    for c in ["CS-LAT", "ROMAN", "NATIVE", "MIXED"]:
        for u, lab in [("prompt", "prompt (500)"), ("proposition", "proposition (100)"), ("topic", "topic (20)")]:
            x = g[c][u]
            rows.append([f"{c} vs. EN-MATCH" if u == "prompt" else "", lab, neg(f"{x['beta']:.3f}"),
                         f"{x['se_robust']:.3f}", fp(x["p"])])
    for u, lab in [("prompt", "prompt (200)"), ("proposition", "proposition (100)"), ("topic", "topic (20)")]:
        x = g["MIXED_vs_CS-LAT"][u]
        rows.append(["MIXED vs. CS-LAT" if u == "prompt" else "", lab, neg(f"{x['beta']:.3f}"),
                     f"{x['se_robust']:.3f}", fp(x["p"])])
    fill_table(T[2], rows, header=["Contrast", "Clustering unit (#)", "β̂", "SE", "p"])
    # move caption + GEE table to Appendix B, after the pairwise table's full-width section
    anchor = P[103]._element
    anchor.addnext(T[2]._tbl)
    anchor.addnext(P[56]._element)

    # drop the per-language, model and descriptor tables (content is in the text / figures)
    for t in (T[7], T[8], T[9]):
        t._tbl.getparent().remove(t._tbl)
    for i in (104, 105, 107, 106, 108):
        delete_par(P[i])

    # ---------------- references ----------------
    last = P[137]
    for k, ref in enumerate(NEW_REFS, start=23):
        newp = insert_after(last, last)
        add_runs(newp, f"[{k}]\t" + ref)
        last = newp

    # ---------------- figures ----------------
    media = {"rId21": "fig1.png", "rId30": "fig2.png", "rId33": "fig3.png", "rId38": "fig4.png", "rId44": "fig5.png"}
    for rel in d.part.rels.values():
        if rel.reltype.endswith("/image") and rel.rId in media:
            rel.target_part._blob = (ROOT / "paper_v2/output" / media[rel.rId]).read_bytes()

    for para in d.paragraphs:
        if para.text.startswith("TABLE "):
            para.paragraph_format.keep_with_next = True
    d.core_properties.title = TEXT["title"]
    OUTDOC.parent.mkdir(parents=True, exist_ok=True)
    d.save(OUTDOC)
    base._fix_custom_abstract(OUTDOC, TEXT["abstract"])
    print("wrote", OUTDOC)


if __name__ == "__main__":
    main()
